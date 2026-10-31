"""
选品池导入：把用户手上的 Excel / CSV 读成统一的候选商品行
==========================================================

输入形态有三种，必须都能吃：
  1. **只有链接**（竞品店铺导出）——产品名称 / 产品链接 / 产品价格(R) / 评分数量
  2. **只有 TSIN**（Takealot 销售报表）——Product Title / TSIN / Qty / Net Sales Amount
  3. **链接 + 销售数据**（自己整理的选品表）

因此这里不做「必须有哪些列」的硬校验，而是**按列名猜字段**，猜不到就留空，
由后续 enrich 步骤按 PLID 补齐。列名同时兼容中英文和常见别名。

为什么不直接用 openpyxl 的默认行为
----------------------------------
`load_workbook` 默认取第一个 sheet。实测销售报表的 sheet 顺序是
`['Sheet3', 'Sheet4', 'Sheet1']`，真实数据在**最后一个** Sheet1；取第一个会
读到一张空表（617 行 vs 13,984 行）。所以这里遍历所有 sheet，按
「识别到的字段数 + 行数」打分挑最像的那张。
"""

from __future__ import annotations

import csv
import io
import re
from typing import Any, Dict, Iterable, List, Optional, Tuple

from services.takealot_public_api import parse_plid

# 字段 -> 列名别名。顺序即优先级：同一列同时像两个字段时归给靠前的那个
# （例："评分数量" 既像 rating 又像 review_count，必须归 review_count）。
FIELD_ALIASES: List[Tuple[str, Tuple[str, ...]]] = [
    ("purchase_url", ("1688链接", "货源链接", "采购链接", "supplier url", "purchase url")),
    ("takealot_url", ("产品链接", "商品链接", "链接", "takealot链接", "url", "link")),
    ("tsin", ("tsin", "tsin号")),
    ("plid", ("plid", "productline", "product_line")),
    # 评论数必须排在评分之前，否则 "评分数量" 会被 "评分" 抢走
    ("review_count", ("评分数量", "评价数量", "评论数量", "评论数", "评价数",
                      "review count", "reviews", "review")),
    ("rating_value", ("产品评分", "评分", "星级", "rating", "star")),
    # Supplier columns must precede generic link/price aliases.
    ("unit_price_cny", ("供货价", "1688价格", "1688单价", "采购单价", "unit price", "wholesale")),
    ("purchase_cost_cny", ("采购成本", "采购价", "进货价")),
    ("purchase_shipping_cny", ("国内运费", "国内快递", "domestic shipping")),
    ("purchase_quantity", ("起订量", "最小起订", "moq", "min order")),
    ("actual_sale_price_zar", ("产品价格", "售价", "销售价", "价格", "price")),
    ("sales_qty", ("销量", "销售数量", "已售", "qty", "quantity", "units", "sold")),
    ("net_sales_zar", ("net sale", "net sales", "净销售", "净额")),
    ("gross_sales_zar", ("gross sale", "gross sales", "总销售", "销售额")),
    ("chinese_product_name", ("中文品名", "中文名称", "中文名", "chinese")),
    ("product_name", ("产品名称", "商品名称", "产品标题", "商品标题", "标题",
                      "product title", "product name", "title", "name")),
    ("note", ("备注", "畅销变体", "note", "remark")),
]

# 至少要认出这些字段里的一个，才算「这张表有货」。
# 只认到 note 或价格不算——那种表没法定位商品。
IDENTITY_FIELDS = ("takealot_url", "tsin", "plid", "product_name")

_NUM_RE = re.compile(r"-?\d+(?:[.,]\d+)?")


def _norm(text: Any) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip().lower()


def detect_columns(headers: Iterable[Any]) -> Dict[str, int]:
    """把表头映射成 {字段: 列下标}。认不出的列直接忽略。"""
    normed = [_norm(h) for h in headers]
    mapping: Dict[str, int] = {}
    used: set = set()
    for field, aliases in FIELD_ALIASES:
        for idx, header in enumerate(normed):
            if not header or idx in used:
                continue
            if any(alias in header for alias in aliases):
                mapping[field] = idx
                used.add(idx)
                break
    return mapping


def to_number(value: Any) -> Optional[float]:
    """把 "R 1,234.56" / "1 234,56" / 325 / "325.0" 都变成数值。

    报表里的价格列常带货币符号、千分位、括号备注（"R325 (含税)"）。逗号既可能
    是千分位也可能是小数点（欧洲写法），所以要靠上下文判断，不能一律删掉——
    "1 234,56" 一律删逗号会变成 123456，价格直接错 100 倍。
    """
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).replace("\u00a0", "").replace(" ", "").strip()
    if not text:
        return None

    has_comma, has_dot = "," in text, "." in text
    if has_comma and has_dot:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")   # 1.234,56
        else:
            text = text.replace(",", "")                      # 1,234.56
    elif has_comma:
        head, _, tail = text.rpartition(",")
        if tail.isdigit() and len(tail) <= 2 and head.replace("-", "").isdigit():
            text = f"{head}.{tail}"                           # 1234,56
        else:
            text = text.replace(",", "")                      # 1,234

    match = _NUM_RE.search(text)
    if match is None:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def to_int(value: Any) -> Optional[int]:
    number = to_number(value)
    return int(number) if number is not None else None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _score_sheet(headers: List[Any], row_count: int) -> int:
    """给 sheet 打分：认出的身份字段权重最高，其次总字段数，最后行数。

    身份字段是硬需求——一张只认出「价格」的表没法定位商品，行数再多也没用。
    """
    mapping = detect_columns(headers)
    identity_hits = sum(1 for f in IDENTITY_FIELDS if f in mapping)
    return identity_hits * 1_000_000 + len(mapping) * 1_000 + min(row_count, 999)


def _rows_from_matrix(matrix: List[List[Any]]) -> Tuple[List[Dict[str, Any]], Dict[str, int], List[str]]:
    """二维表 -> (行字典列表, 列映射, 未识别的列名)。"""
    if not matrix:
        return [], {}, []
    headers = matrix[0]
    mapping = detect_columns(headers)
    for field in ("unit_price_cny", "purchase_cost_cny", "purchase_shipping_cny"):
        if field in mapping and re.search(r"\b(?:usd|eur|gbp|zar|rand)\b|美元|欧元|英镑|兰特", _norm(headers[mapping[field]])):
            raise ValueError(f"{headers[mapping[field]]}不是人民币金额，不能作为人民币采购成本导入")
    unknown = [
        _cell_text(h) for idx, h in enumerate(headers)
        if idx not in mapping.values() and _cell_text(h)
    ]
    rows: List[Dict[str, Any]] = []
    for raw in matrix[1:]:
        if not any(_cell_text(c) for c in raw):
            continue
        item: Dict[str, Any] = {}
        for field, idx in mapping.items():
            if idx < len(raw):
                item[field] = raw[idx]
        if any(_cell_text(item.get(f)) for f in IDENTITY_FIELDS):
            rows.append(item)
    return rows, mapping, unknown


def _read_xlsx(content: bytes) -> Tuple[List[Dict[str, Any]], Dict[str, int], List[str], str]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    best: Optional[Tuple[int, str, List[List[Any]]]] = None
    for sheet in workbook.worksheets:
        matrix: List[List[Any]] = []
        for row in sheet.iter_rows(values_only=True):
            matrix.append(list(row))
            if len(matrix) > 20_000:   # 防御：异常大的表不整张读进内存
                break
        if len(matrix) < 2:
            continue
        score = _score_sheet(matrix[0], len(matrix) - 1)
        if best is None or score > best[0]:
            best = (score, sheet.title, matrix)
    workbook.close()
    if best is None:
        return [], {}, [], ""
    rows, mapping, unknown = _rows_from_matrix(best[2])
    return rows, mapping, unknown, best[1]


def _read_csv(content: bytes) -> Tuple[List[Dict[str, Any]], Dict[str, int], List[str], str]:
    text = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        return [], {}, [], ""
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    matrix = [row for row in csv.reader(io.StringIO(text), dialect)]
    rows, mapping, unknown = _rows_from_matrix(matrix)
    return rows, mapping, unknown, "csv"


def read_table(filename: str, content: bytes) -> Dict[str, Any]:
    """读入上传的表。

    返回：
      {
        "rows": [{字段: 原始值}],       # 只含识别到的列
        "mapping": {字段: 列下标},
        "unknown_headers": [列名],      # 没认出来的列，回给前端提示
        "sheet": "Sheet1",
        "total_rows": 1234,
      }
    """
    lowered = (filename or "").lower()
    if lowered.endswith((".xlsx", ".xlsm")):
        rows, mapping, unknown, sheet = _read_xlsx(content)
    elif lowered.endswith((".csv", ".txt")):
        rows, mapping, unknown, sheet = _read_csv(content)
    else:
        raise ValueError("只支持 .xlsx / .xlsm / .csv；.xls（老二进制格式）请先另存为 xlsx")
    return {
        "rows": rows,
        "mapping": mapping,
        "unknown_headers": unknown,
        "sheet": sheet,
        "total_rows": len(rows),
    }


def normalize_row(raw: Dict[str, Any]) -> Dict[str, Any]:
    """把一行原始值整理成落库字段。不做网络请求。"""
    url = _cell_text(raw.get("takealot_url"))
    plid = parse_plid(raw.get("plid")) or parse_plid(url)
    tsin = _cell_text(raw.get("tsin"))
    return {
        "plid": str(plid) if plid else "",
        "tsin": tsin,
        "takealot_url": url,
        "product_name": _cell_text(raw.get("product_name")),
        "chinese_product_name": _cell_text(raw.get("chinese_product_name")),
        "actual_sale_price_zar": to_number(raw.get("actual_sale_price_zar")),
        "rating_value": to_number(raw.get("rating_value")),
        "review_count": to_int(raw.get("review_count")),
        "sales_qty": to_int(raw.get("sales_qty")),
        "net_sales_zar": to_number(raw.get("net_sales_zar")),
        "gross_sales_zar": to_number(raw.get("gross_sales_zar")),
        "purchase_url": _cell_text(raw.get("purchase_url")),
        "unit_price_cny": to_number(raw.get("unit_price_cny")),
        "purchase_cost_cny": to_number(raw.get("purchase_cost_cny")),
        "purchase_shipping_cny": to_number(raw.get("purchase_shipping_cny")),
        "purchase_quantity": to_int(raw.get("purchase_quantity")),
        "note": _cell_text(raw.get("note")),
    }


def build_takealot_url(plid: Any, slug: str = "") -> str:
    """没有链接时按 PLID 拼一条可点击的链接。

    Takealot 只认 `/slug/PLID{id}`，slug 为空时也能跳转（会 302 到正式 slug），
    所以宁可拼一条粗糙的，也不留空——留空的话前端那张卡片点不动。
    """
    resolved = parse_plid(plid)
    if resolved is None:
        return ""
    if slug:
        return f"https://www.takealot.com/{slug}/PLID{resolved}"
    return f"https://www.takealot.com/p/PLID{resolved}"
